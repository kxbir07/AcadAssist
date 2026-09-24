"""Word Document (DOCX) parser preserving headings and structural sections."""

import io
from pathlib import Path
import docx

from app.services.parsers.base import BaseParser, ExtractedDocument, ExtractedUnit


class DOCXParser(BaseParser):
    """Parser for DOCX documents extracting headings, sections, and paragraphs."""

    def parse(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        doc_title = Path(filename).stem
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)

        if doc.core_properties and doc.core_properties.title:
            prop_title = doc.core_properties.title.strip()
            if prop_title:
                doc_title = prop_title

        units: list[ExtractedUnit] = []
        current_section = "Introduction"
        current_paragraphs: list[str] = []
        section_idx = 1

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else ""

            # Check if this paragraph is a title or heading
            if "heading" in style_name or "title" in style_name:
                if current_paragraphs:
                    unit_text = f"{current_section}\n" + "\n".join(current_paragraphs) if current_section != "Introduction" else "\n".join(current_paragraphs)
                    units.append(
                        ExtractedUnit(
                            content=unit_text.strip(),
                            page_number=section_idx,
                            slide_number=None,
                            title=doc_title,
                            section_title=current_section,
                            content_type="section",
                        )
                    )
                    section_idx += 1
                    current_paragraphs = []
                current_section = text
            else:
                current_paragraphs.append(text)

        # Append any remaining paragraphs
        if current_paragraphs or current_section != "Introduction":
            unit_text = f"{current_section}\n" + "\n".join(current_paragraphs) if current_section != "Introduction" else "\n".join(current_paragraphs)
            if unit_text.strip():
                units.append(
                    ExtractedUnit(
                        content=unit_text.strip(),
                        page_number=section_idx,
                        slide_number=None,
                        title=doc_title,
                        section_title=current_section,
                        content_type="section",
                    )
                )

        # Extract tables
        for table_idx, table in enumerate(doc.tables):
            table_lines: list[str] = []
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    table_lines.append(" | ".join(cells))
            if table_lines:
                units.append(
                    ExtractedUnit(
                        content="\n".join(table_lines),
                        page_number=section_idx + table_idx + 1,
                        slide_number=None,
                        title=doc_title,
                        section_title=f"Table {table_idx + 1}",
                        content_type="table",
                    )
                )

        return ExtractedDocument(
            title=doc_title,
            file_type="docx",
            units=units,
            raw_metadata={"paragraph_count": len(doc.paragraphs), "table_count": len(doc.tables)},
        )
