"""Unit tests for document parsers (PDF, PPT, PPTX, DOCX, TXT)."""

import io
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from app.services.parsers.docx import DOCXParser
from app.services.parsers.factory import ParserFactory
from app.services.parsers.pdf import PDFParser
from app.services.parsers.pptx import PPTXParser
from app.services.parsers.txt import TXTParser


def _create_sample_pdf() -> bytes:
    """Helper to generate a multi-page PDF in-memory."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    # Page 1
    c.drawString(100, 750, "Chapter 1: Operating Systems Overview")
    c.drawString(100, 730, "An operating system manages computer hardware and software resources.")
    c.showPage()
    # Page 2
    c.drawString(100, 750, "Chapter 2: Deadlock Prevention and Avoidance")
    c.drawString(100, 730, "Deadlock occurs when processes hold resources and wait for others.")
    c.drawString(100, 710, "Banker's algorithm guarantees safe state execution.")
    c.showPage()
    c.save()
    return buf.getvalue()


def _create_sample_pptx() -> bytes:
    """Helper to generate a multi-slide PPTX in-memory."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]  # Title slide
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Introduction to Computer Networks"
    slide1.placeholders[1].text = "OSI Model and TCP/IP Architecture"

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])  # Content slide
    slide2.shapes.title.text = "Network Layer Routing"
    tf = slide2.placeholders[1].text_frame
    tf.text = "Routing algorithms: Dijkstra and Bellman-Ford"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _create_sample_docx() -> bytes:
    """Helper to generate a DOCX file in-memory."""
    doc = DocxDocument()
    doc.add_heading("Database Normalization", level=1)
    doc.add_paragraph("First Normal Form (1NF) eliminates repeating groups.")
    doc.add_heading("Second Normal Form (2NF)", level=2)
    doc.add_paragraph("Eliminates partial dependency on candidate keys.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Normal Form"
    table.cell(0, 1).text = "Violation Prevented"
    table.cell(1, 0).text = "3NF"
    table.cell(1, 1).text = "Transitive Dependency"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_parser_page_preservation():
    pdf_bytes = _create_sample_pdf()
    parser = PDFParser()
    extracted = parser.parse(pdf_bytes, "os_concepts.pdf")

    assert extracted.file_type == "pdf"
    assert len(extracted.units) == 2

    # Check page 1
    unit1 = extracted.units[0]
    assert unit1.page_number == 1
    assert "Operating Systems Overview" in unit1.content
    assert unit1.content_type == "page"

    # Check page 2
    unit2 = extracted.units[1]
    assert unit2.page_number == 2
    assert "Deadlock Prevention" in unit2.content
    assert "Banker's algorithm" in unit2.content


def test_pptx_parser_slide_preservation():
    pptx_bytes = _create_sample_pptx()
    parser = PPTXParser()
    extracted = parser.parse(pptx_bytes, "networks.pptx")

    assert extracted.file_type == "pptx"
    assert len(extracted.units) == 2

    slide1 = extracted.units[0]
    assert slide1.slide_number == 1
    assert "Computer Networks" in slide1.content
    assert slide1.content_type == "slide"

    slide2 = extracted.units[1]
    assert slide2.slide_number == 2
    assert "Routing" in slide2.content
    assert "Dijkstra" in slide2.content


def test_docx_parser_structure():
    docx_bytes = _create_sample_docx()
    parser = DOCXParser()
    extracted = parser.parse(docx_bytes, "db_norm.docx")

    assert extracted.file_type == "docx"
    assert len(extracted.units) >= 2

    contents = " ".join(u.content for u in extracted.units)
    assert "First Normal Form" in contents
    assert "Second Normal Form" in contents
    assert "Transitive Dependency" in contents


def test_txt_parser_blocks():
    text_content = (
        "Linear Algebra Lecture 1\nVectors and matrices are fundamental.\n\n"
        "Eigenvalues and Eigenvectors\nCharacteristic equation det(A - lambda*I) = 0."
    )
    parser = TXTParser()
    extracted = parser.parse(text_content.encode("utf-8"), "math.txt")

    assert extracted.file_type == "txt"
    assert len(extracted.units) == 2
    assert "Vectors and matrices" in extracted.units[0].content
    assert "Eigenvalues" in extracted.units[1].content


def test_parser_factory():
    assert isinstance(ParserFactory.get_parser("doc.pdf"), PDFParser)
    assert isinstance(ParserFactory.get_parser("slides.pptx"), PPTXParser)
    assert isinstance(ParserFactory.get_parser("old_slides.ppt"), PPTXParser)
    assert isinstance(ParserFactory.get_parser("notes.docx"), DOCXParser)
    assert isinstance(ParserFactory.get_parser("data.txt"), TXTParser)

    with pytest.raises(ValueError, match="Unsupported document format"):
        ParserFactory.get_parser("script.exe")
